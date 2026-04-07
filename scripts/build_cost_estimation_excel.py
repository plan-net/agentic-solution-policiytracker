#!/usr/bin/env python3
"""
Build Excel Cost Estimation Model from CSV Templates

This script creates a fully-functional Excel workbook from the CSV templates,
including formulas, formatting, data validation, and charts.

Usage:
    python scripts/build_cost_estimation_excel.py

Output:
    docs/cost-estimation/Ingestion_Cost_Estimation_Model.xlsx
"""

import csv
import os
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.chart import PieChart, LineChart, BarChart, Reference
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    print("❌ openpyxl is required. Installing...")
    import subprocess
    subprocess.check_call(["pip", "install", "openpyxl"])
    from openpyxl import Workbook
    from openpyxl.chart import PieChart, LineChart, BarChart, Reference
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation


# Configuration
BASE_DIR = Path(__file__).parent.parent
CSV_DIR = BASE_DIR / "docs" / "cost-estimation" / "csv-templates"
OUTPUT_FILE = BASE_DIR / "docs" / "cost-estimation" / "Ingestion_Cost_Estimation_Model.xlsx"

# Color scheme
COLORS = {
    'header': 'FF4472C4',        # Dark blue
    'input': 'FFD9E1F2',         # Light blue
    'calculated': 'FFF2F2F2',    # Light gray
    'total': 'FFE2EFDA',         # Light green
    'warning': 'FFFFC000',       # Orange
    'alert': 'FFFCE4D6',         # Light red
}


def read_csv_file(csv_path: Path) -> list:
    """Read CSV file and return rows."""
    with open(csv_path, 'r', encoding='utf-8') as f:
        reader = csv.reader(f)
        return list(reader)


def create_workbook() -> Workbook:
    """Create a new Excel workbook."""
    wb = Workbook()
    # Remove default sheet
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
    return wb


def import_csv_to_sheet(wb: Workbook, csv_path: Path, sheet_name: str):
    """Import CSV data to a new sheet."""
    print(f"  📄 Importing {csv_path.name} → {sheet_name}")

    rows = read_csv_file(csv_path)
    ws = wb.create_sheet(title=sheet_name)

    # Write data
    for row_idx, row_data in enumerate(rows, start=1):
        for col_idx, cell_value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)

            # Try to detect and convert formulas
            if isinstance(cell_value, str) and cell_value.startswith('='):
                cell.value = cell_value  # Excel will interpret as formula
            # Replace [DATE] placeholders
            elif '[DATE]' in str(cell_value):
                cell.value = cell_value.replace('[DATE]', datetime.now().strftime('%Y-%m-%d'))
            else:
                cell.value = cell_value

    return ws


def apply_basic_formatting(ws):
    """Apply basic formatting to a worksheet."""
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass

        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width


def format_dashboard(ws):
    """Apply specific formatting to Dashboard sheet."""
    print("  🎨 Formatting Dashboard sheet")

    # Header formatting
    header_font = Font(bold=True, size=14, color='FFFFFFFF')
    header_fill = PatternFill(start_color=COLORS['header'], end_color=COLORS['header'], fill_type='solid')

    # Format title
    ws['A1'].font = Font(bold=True, size=16)

    # Format section headers (look for cells with "QUICK COST ESTIMATOR", etc.)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                if cell.value.isupper() and len(cell.value) > 10:
                    cell.font = header_font
                    cell.fill = header_fill

    apply_basic_formatting(ws)


def format_input_parameters(ws):
    """Apply specific formatting to Input Parameters sheet."""
    print("  🎨 Formatting Input Parameters sheet")

    # Add data validation for model selection
    # Find cells B18 and B19 (LLM and Embedding model)

    # LLM Model dropdown (B18)
    llm_validation = DataValidation(
        type="list",
        formula1='"gpt-4o-mini,gpt-4o,claude-sonnet-4,claude-haiku"',
        allow_blank=False
    )
    ws.add_data_validation(llm_validation)
    llm_validation.add('B18')

    # Embedding Model dropdown (B19)
    embed_validation = DataValidation(
        type="list",
        formula1='"text-embedding-3-small,text-embedding-3-large,text-embedding-ada-002"',
        allow_blank=False
    )
    ws.add_data_validation(embed_validation)
    embed_validation.add('B19')

    # Prompt Caching dropdown (B20)
    cache_validation = DataValidation(
        type="list",
        formula1='"YES,NO"',
        allow_blank=False
    )
    ws.add_data_validation(cache_validation)
    cache_validation.add('B20')

    # Chunk Limiting dropdown (B26)
    chunk_validation = DataValidation(
        type="list",
        formula1='"YES,NO"',
        allow_blank=False
    )
    ws.add_data_validation(chunk_validation)
    chunk_validation.add('B26')

    # Highlight input cells
    input_fill = PatternFill(start_color=COLORS['input'], end_color=COLORS['input'], fill_type='solid')
    for cell_ref in ['B4', 'B5', 'B6', 'B18', 'B19', 'B20']:
        if cell_ref in ws:
            ws[cell_ref].fill = input_fill

    apply_basic_formatting(ws)


def format_pricing_reference(ws):
    """Apply specific formatting to Pricing Reference sheet."""
    print("  🎨 Formatting Pricing Reference sheet")

    # Highlight cells that need updating (Exa.ai and Apify costs)
    warning_fill = PatternFill(start_color=COLORS['warning'], end_color=COLORS['warning'], fill_type='solid')
    warning_font = Font(bold=True)

    # Mark C7 and C8 as needing updates
    for cell_ref in ['C7', 'C8']:
        if cell_ref in ws:
            ws[cell_ref].fill = warning_fill
            ws[cell_ref].font = warning_font
            # Add comment
            ws[cell_ref].comment = "UPDATE REQUIRED - Add actual API pricing"

    apply_basic_formatting(ws)


def format_cost_calculator(ws):
    """Apply specific formatting to Cost Calculator sheet."""
    print("  🎨 Formatting Cost Calculator sheet")

    # Highlight total cells
    total_fill = PatternFill(start_color=COLORS['total'], end_color=COLORS['total'], fill_type='solid')
    total_font = Font(bold=True, size=11)

    # Key total cells
    total_cells = ['B15', 'B43', 'B61', 'B71', 'B80']
    for cell_ref in total_cells:
        if cell_ref in ws:
            ws[cell_ref].fill = total_fill
            ws[cell_ref].font = total_font

    apply_basic_formatting(ws)


def format_scenario_comparison(ws):
    """Apply specific formatting to Scenario Comparison sheet."""
    print("  🎨 Formatting Scenario Comparison sheet")
    apply_basic_formatting(ws)


def format_historical_validation(ws):
    """Apply specific formatting to Historical Validation sheet."""
    print("  🎨 Formatting Historical Validation sheet")
    apply_basic_formatting(ws)


def add_dashboard_charts(ws):
    """Add charts to Dashboard sheet."""
    print("  📊 Adding charts to Dashboard")

    # Note: Chart creation requires exact cell ranges
    # Since CSV structure might vary, we'll create simple charts

    try:
        # Pie chart for cost breakdown
        pie = PieChart()
        pie.title = "Monthly Cost Breakdown"
        pie.style = 10

        # Data range (Component names and costs) - adjust based on actual structure
        labels = Reference(ws, min_col=1, min_row=20, max_row=23)
        data = Reference(ws, min_col=2, min_row=20, max_row=23)

        pie.add_data(data, titles_from_data=False)
        pie.set_categories(labels)

        ws.add_chart(pie, "E20")

    except Exception as e:
        print(f"    ⚠️  Could not create charts: {e}")


def create_excel_from_csvs():
    """Main function to create Excel file from CSV templates."""
    print("\n🚀 Building Cost Estimation Excel Model\n")

    # Check if CSV directory exists
    if not CSV_DIR.exists():
        print(f"❌ CSV templates directory not found: {CSV_DIR}")
        return False

    # CSV files to import (in order)
    csv_files = [
        ('01_Dashboard.csv', 'Dashboard'),
        ('02_Input_Parameters.csv', 'Input Parameters'),
        ('03_Pricing_Reference.csv', 'Pricing Reference'),
        ('04_Cost_Calculator.csv', 'Cost Calculator'),
        ('05_Scenario_Comparison.csv', 'Scenario Comparison'),
        ('06_Historical_Validation.csv', 'Historical Validation'),
    ]

    # Create workbook
    print("📋 Creating workbook...")
    wb = create_workbook()

    # Import and format each sheet
    for csv_filename, sheet_name in csv_files:
        csv_path = CSV_DIR / csv_filename

        if not csv_path.exists():
            print(f"  ⚠️  CSV not found: {csv_filename}")
            continue

        # Import CSV
        ws = import_csv_to_sheet(wb, csv_path, sheet_name)

        # Apply sheet-specific formatting
        if sheet_name == 'Dashboard':
            format_dashboard(ws)
            add_dashboard_charts(ws)
        elif sheet_name == 'Input Parameters':
            format_input_parameters(ws)
        elif sheet_name == 'Pricing Reference':
            format_pricing_reference(ws)
        elif sheet_name == 'Cost Calculator':
            format_cost_calculator(ws)
        elif sheet_name == 'Scenario Comparison':
            format_scenario_comparison(ws)
        elif sheet_name == 'Historical Validation':
            format_historical_validation(ws)

    # Set active sheet to Dashboard
    wb.active = wb['Dashboard']

    # Save workbook
    print(f"\n💾 Saving workbook to: {OUTPUT_FILE}")
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)

    print(f"\n✅ Success! Excel file created: {OUTPUT_FILE}")
    print(f"\n📂 Location: {OUTPUT_FILE.relative_to(BASE_DIR)}")
    print("\n🎯 Next steps:")
    print("   1. Open the Excel file")
    print("   2. Update Pricing Reference sheet cells C7 and C8 with actual API costs")
    print("   3. Configure your parameters in Input Parameters sheet")
    print("   4. Review Dashboard for cost estimates")
    print("\n📖 For detailed instructions, see:")
    print("   docs/cost-estimation/EXCEL_ASSEMBLY_GUIDE.md")

    return True


if __name__ == "__main__":
    try:
        success = create_excel_from_csvs()
        exit(0 if success else 1)
    except Exception as e:
        print(f"\n❌ Error: {e}")
        import traceback
        traceback.print_exc()
        exit(1)
